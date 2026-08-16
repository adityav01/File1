from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import datetime
from typing import Iterable
from xml.etree import ElementTree as ET

import httpx

SOAP_NS = "http://tempuri.org/"
SOAP_ENV = """<?xml version="1.0" encoding="utf-8"?>
<soap:Envelope xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"
               xmlns:xsd="http://www.w3.org/2001/XMLSchema"
               xmlns:soap="http://schemas.xmlsoap.org/soap/envelope/">
  <soap:Body>
    <GetTransactionsLog xmlns="http://tempuri.org/">
      <FromDateTime>{from_time}</FromDateTime>
      <ToDateTime>{to_time}</ToDateTime>
      <SerialNumber>{serial}</SerialNumber>
      <UserName>{username}</UserName>
      <UserPassword>{password}</UserPassword>
      <strDataList></strDataList>
    </GetTransactionsLog>
  </soap:Body>
</soap:Envelope>
"""


@dataclass(frozen=True)
class PunchRecord:
    biometric_user_id: str
    punch_time: datetime
    direction: str
    device_serial: str = ""
    raw: str = ""


def normalize_direction(value: str | None) -> str:
    if not value:
        return "UNKNOWN"
    token = value.strip().upper()
    mapping = {
        "IN": "IN",
        "I": "IN",
        "0": "IN",
        "CHECK-IN": "IN",
        "CHECKIN": "IN",
        "ENTRY": "IN",
        "OUT": "OUT",
        "O": "OUT",
        "1": "OUT",
        "CHECK-OUT": "OUT",
        "CHECKOUT": "OUT",
        "EXIT": "OUT",
    }
    return mapping.get(token, "UNKNOWN")


def parse_punch_time(value: str) -> datetime:
    text = value.strip().strip('"')
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d-%m-%Y %H:%M:%S",
        "%d/%m/%Y %H:%M:%S",
        "%d-%m-%Y %H:%M",
        "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M:%S%z",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise ValueError(f"Unsupported punch time: {value}")


def parse_log_line(line: str) -> PunchRecord | None:
    text = line.strip()
    if not text:
        return None
    if "\t" in text:
        parts = [part.strip() for part in text.split("\t")]
    elif ";" in text:
        parts = [part.strip() for part in text.split(";")]
    else:
        parts = [part.strip() for part in re.split(r"\s*,\s*", text)]

    if len(parts) < 2:
        return None

    # eBioServer device logs: LogId, DateTime, EmpCode, Device, Location, Direction
    if len(parts) >= 6 and _looks_like_datetime(parts[1]):
        return PunchRecord(
            biometric_user_id=parts[2],
            punch_time=parse_punch_time(parts[1]),
            direction=normalize_direction(parts[5]),
            device_serial=parts[3],
            raw=text,
        )

    # eTimeTrackLite GetTransactionsLog: UserId, DateTime, Direction, Serial
    if _looks_like_datetime(parts[1]):
        return PunchRecord(
            biometric_user_id=parts[0],
            punch_time=parse_punch_time(parts[1]),
            direction=normalize_direction(parts[2] if len(parts) > 2 else ""),
            device_serial=parts[3] if len(parts) > 3 else "",
            raw=text,
        )
    return None


def _looks_like_datetime(value: str) -> bool:
    try:
        parse_punch_time(value)
        return True
    except ValueError:
        return False


def parse_str_data_list(payload: str) -> list[PunchRecord]:
    records: list[PunchRecord] = []
    for line in payload.splitlines():
        record = parse_log_line(line)
        if record:
            records.append(record)
    return records


def parse_soap_response(xml_text: str) -> list[PunchRecord]:
    root = ET.fromstring(xml_text)
    node = root.find(f".//{{{SOAP_NS}}}strDataList")
    if node is None or not node.text:
        node = root.find(".//{*}strDataList")
    if node is None or not node.text:
        return []
    return parse_str_data_list(node.text)


def _norm_key(key: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(key).lower())


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value).strip()


def _first(row: dict[str, str], *keys: str) -> str:
    for key in keys:
        if row.get(key):
            return row[key]
    return ""


def punch_from_row(row: dict[str, str]) -> PunchRecord | None:
    user_id = _first(
        row,
        "biometricuserid",
        "userid",
        "userno",
        "empcode",
        "employeecode",
        "enrollnumber",
        "enrollno",
    )
    punch_time = _first(row, "punchtime", "logtime", "datetime", "dateandtime", "punchdatetime")
    if not punch_time and row.get("date"):
        date_part = row["date"].split()[0]
        time_part = row.get("time") or "00:00:00"
        punch_time = f"{date_part} {time_part}"
    if not user_id or not punch_time:
        return None
    return PunchRecord(
        biometric_user_id=user_id,
        punch_time=parse_punch_time(punch_time),
        direction=normalize_direction(_first(row, "direction", "inout", "status", "type")),
        device_serial=_first(row, "deviceserial", "serial", "devicename"),
        raw=",".join(row.values()),
    )


def employee_from_row(row: dict[str, str]) -> dict[str, str] | None:
    emp_code = _first(row, "empcode", "employeecode", "code")
    name = _first(row, "name", "employeename")
    user_id = _first(row, "biometricuserid", "userid", "userno", "enrollnumber")
    if not emp_code or not name or not user_id:
        return None
    return {
        "emp_code": emp_code,
        "name": name,
        "biometric_user_id": user_id,
        "department": _first(row, "department", "dept") or "Operations",
        "designation": _first(row, "designation", "role") or "Staff",
        "basic": _first(row, "basic") or "0",
        "hra": _first(row, "hra") or "0",
        "other_allowance": _first(row, "otherallowance", "other", "allowance") or "0",
    }


def _rows_from_csv(content: str) -> list[dict[str, str]]:
    sample = content[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;") if sample.strip() else csv.excel
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(content), dialect=dialect)
    if not reader.fieldnames:
        return []
    rows = []
    for raw in reader:
        rows.append({_norm_key(k): _cell_text(v) for k, v in raw.items() if k})
    return rows


def _rows_from_xlsx(data: bytes) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        return []
    keys = [_norm_key(_cell_text(col) or f"col{index}") for index, col in enumerate(header)]
    rows = []
    for values in rows_iter:
        row = {keys[i]: _cell_text(values[i] if i < len(values) else "") for i in range(len(keys))}
        if any(row.values()):
            rows.append(row)
    return rows


def parse_csv(content: str) -> list[PunchRecord]:
    records = [record for row in _rows_from_csv(content) if (record := punch_from_row(row))]
    if records:
        return records
    parsed = []
    lines = content.splitlines()
    start = 1 if lines and "," in lines[0] and not _looks_like_datetime(lines[0].split(",")[1] if "," in lines[0] else "") else 0
    for line in lines[start:]:
        record = parse_log_line(line)
        if record:
            parsed.append(record)
    return parsed


def parse_employees_csv(content: str) -> list[dict[str, str]]:
    return [row for item in _rows_from_csv(content) if (row := employee_from_row(item))]


def parse_upload(filename: str, data: bytes) -> list[PunchRecord]:
    name = filename.lower()
    if name.endswith(".xlsx") or name.endswith(".xls"):
        return [record for row in _rows_from_xlsx(data) if (record := punch_from_row(row))]
    text = data.decode("utf-8-sig")
    return parse_csv(text)


def parse_employee_upload(filename: str, data: bytes) -> list[dict[str, str]]:
    name = filename.lower()
    if name.endswith(".xlsx") or name.endswith(".xls"):
        return [row for item in _rows_from_xlsx(data) if (row := employee_from_row(item))]
    return parse_employees_csv(data.decode("utf-8-sig"))


class EsslSoapClient:
    def __init__(self, soap_url: str, timeout: float = 30.0):
        self.soap_url = soap_url.rstrip("/")
        self.timeout = timeout

    def fetch_transactions(
        self,
        *,
        from_time: datetime,
        to_time: datetime,
        username: str,
        password: str,
        serial_number: str,
    ) -> list[PunchRecord]:
        body = SOAP_ENV.format(
            from_time=from_time.strftime("%Y-%m-%d %H:%M:%S"),
            to_time=to_time.strftime("%Y-%m-%d %H:%M:%S"),
            serial=serial_number,
            username=username,
            password=password,
        )
        headers = {
            "Content-Type": "text/xml; charset=utf-8",
            "SOAPAction": "http://tempuri.org/GetTransactionsLog",
        }
        response = httpx.post(self.soap_url, content=body, headers=headers, timeout=self.timeout)
        response.raise_for_status()
        return parse_soap_response(response.text)


def punches_from_push_payload(payload: dict) -> Iterable[PunchRecord]:
    punch = payload.get("PunchLog") or payload
    user_id = str(punch.get("UserId") or punch.get("biometric_user_id") or "")
    log_time = str(punch.get("LogTime") or punch.get("punch_time") or "")
    if not user_id or not log_time:
        return []
    serial = str(payload.get("SerialNumber") or punch.get("device_serial") or "")
    return [
        PunchRecord(
            biometric_user_id=user_id,
            punch_time=parse_punch_time(log_time.replace("T", " ").split("+")[0]),
            direction=normalize_direction(str(punch.get("Type") or punch.get("direction") or "")),
            device_serial=serial,
            raw=str(payload),
        )
    ]
